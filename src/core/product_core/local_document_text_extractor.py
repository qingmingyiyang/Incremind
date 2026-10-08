from __future__ import annotations

import locale
import subprocess
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol

from .source_content_read import SourceContentReadError


class LocalDocumentTextExtractorError(SourceContentReadError):
    """Raised when a local document text extractor cannot extract authorized content."""


BUILTIN_DOCUMENT_TEXT_COMMAND = "builtin:document-text"


class DocumentTextExtractorPort(Protocol):
    def extract(self, document_path: Path, media_type: str) -> str: ...


@dataclass(frozen=True, slots=True)
class BuiltinDocumentTextExtractor:
    """Extract DOCX/PDF text with libraries already shipped in the sidecar runtime."""

    provider_name: str = "builtin-document-text"

    def extract(self, document_path: Path, media_type: str) -> str:
        clean_path = document_path.expanduser().resolve(strict=False)
        if not clean_path.is_file():
            raise LocalDocumentTextExtractorError("authorized document does not exist")
        clean_media_type = _required_str(media_type, "media_type")
        try:
            if clean_media_type == "application/pdf" or clean_path.suffix.lower() == ".pdf":
                text = _extract_pdf_text(clean_path)
            elif (
                clean_media_type
                == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                or clean_path.suffix.lower() == ".docx"
            ):
                text = _extract_docx_text(clean_path)
            else:
                raise LocalDocumentTextExtractorError("built-in extractor supports DOCX and PDF only")
        except LocalDocumentTextExtractorError:
            raise
        except Exception as error:
            raise LocalDocumentTextExtractorError("built-in document extraction failed") from error
        clean_text = text.strip()
        if not clean_text:
            raise LocalDocumentTextExtractorError("local document text extractor returned no text")
        return clean_text


@dataclass(frozen=True, slots=True)
class LocalCommandDocumentTextExtractor:
    """Document text extractor backed by an explicitly configured local command.

    The command may point to a local extractor such as MarkItDown:
    ("powershell", "-File", "run-markitdown.ps1", "-InputPath", "{document_path}", "-Stdout")
    """

    command: Sequence[str]
    enabled: bool = False
    provider_name: str = "local-command-document-text"
    timeout_seconds: float = 60.0

    def extract(self, document_path: Path, media_type: str) -> str:
        if self.enabled is not True:
            raise LocalDocumentTextExtractorError("local document text extractor is disabled")
        if self.timeout_seconds <= 0:
            raise LocalDocumentTextExtractorError("local document text extractor timeout must be positive")
        clean_media_type = _required_str(media_type, "media_type")
        clean_path = document_path.expanduser().resolve(strict=False)
        if not clean_path.exists():
            raise LocalDocumentTextExtractorError("authorized document does not exist")
        if not clean_path.is_file():
            raise LocalDocumentTextExtractorError("authorized document path is not a file")
        command = _command_for_document(self.command, clean_path, clean_media_type)
        try:
            completed = subprocess.run(
                command,
                check=False,
                capture_output=True,
                timeout=self.timeout_seconds,
            )
        except FileNotFoundError as exc:
            raise LocalDocumentTextExtractorError("local document text extractor executable not found") from exc
        except subprocess.TimeoutExpired as exc:
            raise LocalDocumentTextExtractorError("local document text extractor timed out") from exc
        stdout = _decode_command_output(completed.stdout)
        stderr = _decode_command_output(completed.stderr)
        if completed.returncode != 0:
            detail = (stderr or stdout or "local document text extractor failed").strip()
            raise LocalDocumentTextExtractorError(_preview_error(detail))
        text = stdout.strip()
        if not text:
            raise LocalDocumentTextExtractorError("local document text extractor returned no text")
        return text


def document_text_extractors_for_allowed_documents(
    extractor: DocumentTextExtractorPort,
) -> Mapping[str, Callable[[Path, str], str]]:
    return {
        "application/pdf": extractor.extract,
        "application/msword": extractor.extract,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": extractor.extract,
    }


def _extract_docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append("\t".join(cells))
    return "\n".join(blocks)


def _extract_pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            if reader.decrypt("") == 0:
                raise LocalDocumentTextExtractorError("encrypted PDF requires a password")
        except LocalDocumentTextExtractorError:
            raise
        except Exception as error:
            raise LocalDocumentTextExtractorError("encrypted PDF requires a password") from error
    return "\n".join(text for page in reader.pages if (text := (page.extract_text() or "").strip()))


def _command_for_document(command: Sequence[str], document_path: Path, media_type: str) -> tuple[str, ...]:
    if not isinstance(command, Sequence) or isinstance(command, (str, bytes)) or not command:
        raise LocalDocumentTextExtractorError("local document text extractor command is not configured")
    replaced: list[str] = []
    has_document_placeholder = False
    for part in command:
        if not isinstance(part, str) or not part:
            raise LocalDocumentTextExtractorError(
                "local document text extractor command parts must be non-empty strings"
            )
        if "{document_path}" in part:
            has_document_placeholder = True
        replaced.append(
            part.replace("{document_path}", str(document_path)).replace("{media_type}", media_type)
        )
    if not has_document_placeholder:
        replaced.append(str(document_path))
    return tuple(replaced)


def _preview_error(text: str, limit: int = 240) -> str:
    compact = " ".join(text.split())
    if len(compact) <= limit:
        return compact
    return f"{compact[: limit - 1]}..."


def _decode_command_output(value: bytes | str | None) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    encodings = ("utf-8", locale.getpreferredencoding(False))
    attempted: set[str] = set()
    for encoding in encodings:
        normalized = encoding.lower()
        if normalized in attempted:
            continue
        attempted.add(normalized)
        try:
            return value.decode(encoding)
        except (LookupError, UnicodeDecodeError):
            continue
    return value.decode("utf-8", errors="replace")


def _required_str(value: object, field_name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise LocalDocumentTextExtractorError(f"{field_name} is required")
    return value.strip()
