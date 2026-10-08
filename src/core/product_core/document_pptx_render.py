"""Deterministic editable PPTX rendering for governed Document Delivery."""

from __future__ import annotations

import html
import re
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo


class DocumentPptxRenderError(ValueError):
    """Raised when a frozen Document cannot be represented as bounded slides."""


@dataclass(frozen=True, slots=True)
class DocumentPptxRenderResult:
    content: bytes
    slide_html: str
    title: str
    document_id: str
    revision: int | None
    slide_count: int


@dataclass(frozen=True, slots=True)
class SlideLayoutPlan:
    """Serializable bounded storyboard frozen before PPTX delivery."""

    document_id: str
    title: str
    revision: int | None
    slides: tuple[_TextSlide | _TableSlide, ...]

    def payload(self) -> dict[str, object]:
        return {
            "schema_version": "1.0.0", "document_id": self.document_id,
            "title": self.title, "revision": self.revision,
            "slides": [
                {"kind": "text", "title": slide.title,
                 "lines": [[value, bullet] for value, bullet in slide.lines]}
                if isinstance(slide, _TextSlide) else
                {"kind": "table", "title": slide.title,
                 "rows": [list(row) for row in slide.rows]}
                for slide in self.slides
            ],
        }

    @classmethod
    def from_payload(cls, value: Mapping[str, object]) -> "SlideLayoutPlan":
        if set(value) != {"schema_version", "document_id", "title", "revision", "slides"}:
            raise DocumentPptxRenderError("slide layout plan fields are invalid")
        if value.get("schema_version") != "1.0.0":
            raise DocumentPptxRenderError("slide layout plan version is invalid")
        document_id, title = value.get("document_id"), value.get("title")
        revision, raw_slides = value.get("revision"), value.get("slides")
        if not isinstance(document_id, str) or not document_id or not isinstance(title, str) or not title:
            raise DocumentPptxRenderError("slide layout plan identity is invalid")
        if revision is not None and (isinstance(revision, bool) or not isinstance(revision, int)):
            raise DocumentPptxRenderError("slide layout plan revision is invalid")
        if not isinstance(raw_slides, list) or len(raw_slides) > 79:
            raise DocumentPptxRenderError("slide layout plan slides are invalid")
        slides: list[_TextSlide | _TableSlide] = []
        for item in raw_slides:
            if (
                not isinstance(item, Mapping)
                or not isinstance(item.get("title"), str)
                or not 0 < len(item["title"]) <= 120
            ):
                raise DocumentPptxRenderError("slide layout plan slide is invalid")
            if item.get("kind") == "text" and set(item) == {"kind", "title", "lines"} and isinstance(item.get("lines"), list):
                lines = tuple((line[0], line[1]) for line in item["lines"] if isinstance(line, list) and len(line) == 2 and isinstance(line[0], str) and isinstance(line[1], bool))
                if (
                    len(lines) != len(item["lines"])
                    or len(lines) > 11
                    or any(len(line[0]) > 480 for line in lines)
                ): raise DocumentPptxRenderError("slide layout plan text is invalid")
                slides.append(_TextSlide(item["title"], lines))
            elif item.get("kind") == "table" and set(item) == {"kind", "title", "rows"} and isinstance(item.get("rows"), list):
                rows = tuple(tuple(cell for cell in row) for row in item["rows"] if isinstance(row, list) and all(isinstance(cell, str) for cell in row))
                if (
                    not rows or len(rows) != len(item["rows"])
                    or len(rows) > 12 or len(rows[0]) > 8
                    or not all(len(row) == len(rows[0]) for row in rows)
                    or any(len(cell) > 280 for row in rows for cell in row)
                ): raise DocumentPptxRenderError("slide layout plan table is invalid")
                slides.append(_TableSlide(item["title"], rows))
            else: raise DocumentPptxRenderError("slide layout plan slide is invalid")
        return cls(document_id, title, revision, tuple(slides))


@dataclass(frozen=True, slots=True)
class _TextSlide:
    title: str
    lines: tuple[tuple[str, bool], ...]


@dataclass(frozen=True, slots=True)
class _TableSlide:
    title: str
    rows: tuple[tuple[str, ...], ...]


@dataclass(frozen=True, slots=True)
class DocumentPptxRenderer:
    """Map a conservative Markdown subset to editable PowerPoint shapes."""

    max_slides: int = 80
    max_lines_per_slide: int = 11
    max_chars_per_slide: int = 1_300
    max_table_rows: int = 12
    max_table_columns: int = 8

    def plan(self, document: Mapping[str, object], markdown: str) -> SlideLayoutPlan:
        if not isinstance(document, Mapping) or not isinstance(markdown, str):
            raise DocumentPptxRenderError("document and markdown are required")
        document_id = str(document.get("id") or "")
        title = str(document.get("title") or "").strip()
        revision = document.get("revision")
        if not document_id or not title:
            raise DocumentPptxRenderError("document identity and title are required")
        if isinstance(revision, bool) or not isinstance(revision, int):
            revision = None
        storyboard = _storyboard(
            title,
            markdown,
            max_lines=self.max_lines_per_slide,
            max_chars=self.max_chars_per_slide,
            max_table_rows=self.max_table_rows,
            max_table_columns=self.max_table_columns,
        )
        if len(storyboard) + 1 > self.max_slides:
            raise DocumentPptxRenderError("document exceeds governed PPTX slide budget")
        return SlideLayoutPlan(document_id, title, revision, storyboard)

    def render(
        self, document: Mapping[str, object], markdown: str, *,
        layout_plan: SlideLayoutPlan | None = None,
    ) -> DocumentPptxRenderResult:
        plan = layout_plan or self.plan(document, markdown)
        document_id = str(document.get("id") or "")
        title = str(document.get("title") or "").strip()
        revision = document.get("revision")
        if isinstance(revision, bool) or not isinstance(revision, int): revision = None
        if (plan.document_id, plan.title, plan.revision) != (document_id, title, revision):
            raise DocumentPptxRenderError("slide layout plan identity conflicts")
        storyboard = plan.slides
        content = _render_pptx(title, document_id, revision, storyboard)
        return DocumentPptxRenderResult(
            content,
            _render_slide_html(title, revision, storyboard),
            title,
            document_id,
            revision,
            len(storyboard) + 1,
        )


def _storyboard(
    document_title: str,
    markdown: str,
    *,
    max_lines: int,
    max_chars: int,
    max_table_rows: int,
    max_table_columns: int,
) -> tuple[_TextSlide | _TableSlide, ...]:
    slides: list[_TextSlide | _TableSlide] = []
    section = "内容概览"
    pending: list[tuple[str, bool]] = []
    code = False
    lines = markdown.splitlines()
    index = 0

    def flush_text() -> None:
        nonlocal pending
        if not pending:
            return
        chunks: list[list[tuple[str, bool]]] = []
        chunk: list[tuple[str, bool]] = []
        char_count = 0
        visual_units = 0
        for item in pending:
            item_units = _visual_line_units(item[0], bullet=item[1])
            if item_units > 13:
                raise DocumentPptxRenderError("paragraph exceeds visible PPTX layout budget")
            if chunk and (
                len(chunk) >= max_lines
                or char_count + len(item[0]) > max_chars
                or visual_units + item_units > 13
            ):
                chunks.append(chunk)
                chunk, char_count, visual_units = [], 0, 0
            chunk.append(item)
            char_count += len(item[0])
            visual_units += item_units
        if chunk:
            chunks.append(chunk)
        for part, values in enumerate(chunks, 1):
            suffix = f"（{part}）" if len(chunks) > 1 else ""
            slides.append(_TextSlide(f"{section}{suffix}"[:120], tuple(values)))
        pending = []

    while index < len(lines):
        stripped = lines[index].strip()
        if stripped.startswith("```"):
            code = not code
            index += 1
            continue
        heading = re.fullmatch(r"#{1,3}\s+(.+)", stripped)
        if heading and not code:
            flush_text()
            heading_text = _plain(heading.group(1)) or document_title
            section = heading_text[:120]
            for continuation in _chunks(heading_text[120:], 480):
                pending.append((continuation, False))
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and _is_table_separator(lines[index + 1]):
            flush_text()
            rows: list[tuple[str, ...]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                if not _is_table_separator(lines[index]):
                    row = tuple(_plain(cell) for cell in _table_cells(lines[index]))
                    rows.append(row)
                index += 1
            if not rows:
                continue
            width = max(len(row) for row in rows)
            if width > max_table_columns:
                raise DocumentPptxRenderError("table exceeds governed PPTX column budget")
            header, data_rows = rows[0], rows[1:]
            batches = [data_rows[start : start + max_table_rows - 1]
                       for start in range(0, len(data_rows), max_table_rows - 1)] or [[]]
            for batch_index, data_batch in enumerate(batches, 1):
                batch = [header, *data_batch]
                normalized = tuple(row + ("",) * (width - len(row)) for row in batch)
                _assert_table_fit(normalized)
                suffix = f"（{batch_index}）" if len(batches) > 1 else ""
                slides.append(_TableSlide(f"{section}{suffix}"[:120], normalized))
            continue
        if stripped:
            bullet = re.fullmatch(r"(?:[-+*]|\d+\.)\s+(.+)", stripped)
            value = bullet.group(1) if bullet else stripped.lstrip(">").strip()
            prefix = "代码：" if code else ""
            plain = _plain(value)
            if plain:
                for part_index, part in enumerate(_chunks(prefix + plain, 480)):
                    pending.append((part, bullet is not None and part_index == 0))
        index += 1
    flush_text()
    return tuple(slides)


def _render_pptx(
    title: str,
    document_id: str,
    revision: int | None,
    storyboard: tuple[_TextSlide | _TableSlide, ...],
) -> bytes:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise DocumentPptxRenderError("python-pptx runtime is unavailable") from exc

    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    properties = presentation.core_properties
    properties.title = title
    properties.subject = "Chriptmas OS governed Document Delivery"
    properties.identifier = document_id
    properties.version = str(revision or "")
    fixed = datetime(2000, 1, 1, tzinfo=UTC)
    properties.created = fixed
    properties.modified = fixed

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _background(title_slide, RGBColor(0xF7, 0xF4, 0xEE))
    _textbox(title_slide, title, 0.85, 2.35, 11.65, 1.2, 30, RGBColor(0x91, 0x01, 0x01), bold=True, align=PP_ALIGN.CENTER)
    _textbox(title_slide, f"Document revision {revision or 'unknown'}", 1.4, 3.75, 10.55, 0.45, 12, RGBColor(0x68, 0x60, 0x74), align=PP_ALIGN.CENTER)

    for spec in storyboard:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _background(slide, RGBColor(0xF7, 0xF4, 0xEE))
        _textbox(slide, spec.title, 0.7, 0.45, 11.9, 0.8, 25, RGBColor(0x91, 0x01, 0x01), bold=True)
        if isinstance(spec, _TextSlide):
            shape = slide.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(11.5), Inches(5.25))
            frame = shape.text_frame
            frame.clear()
            frame.word_wrap = True
            for idx, (text, bullet) in enumerate(spec.lines):
                paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
                paragraph.text = text
                paragraph.level = 0
                paragraph.font.name = "Microsoft YaHei"
                paragraph.font.size = Pt(18)
                paragraph.font.color.rgb = RGBColor(0x25, 0x25, 0x25)
                paragraph.space_after = Pt(8)
                if bullet:
                    paragraph.text = f"• {text}"
        else:
            rows, columns = len(spec.rows), len(spec.rows[0])
            table = slide.shapes.add_table(rows, columns, Inches(0.75), Inches(1.55), Inches(11.85), Inches(4.9)).table
            for row_index, row in enumerate(spec.rows):
                for column_index, value in enumerate(row):
                    cell = table.cell(row_index, column_index)
                    cell.text = value
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(0xEE, 0xE8, 0xDE) if row_index == 0 else RGBColor(0xFF, 0xFF, 0xFF)
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.name = "Microsoft YaHei"
                        paragraph.font.size = Pt(13)
                        paragraph.font.bold = row_index == 0
                        paragraph.font.color.rgb = RGBColor(0x25, 0x25, 0x25)

    output = BytesIO()
    presentation.save(output)
    return _canonicalize(output.getvalue())


def _render_slide_html(
    title: str,
    revision: int | None,
    storyboard: tuple[_TextSlide | _TableSlide, ...],
) -> str:
    """Render the same bounded storyboard as self-contained 16:9 print input."""

    def text(value: str) -> str:
        return html.escape(value, quote=True)

    slides = [
        "<section class=\"slide title-slide\"><h1>"
        f"{text(title)}</h1><p>Document revision {revision or 'unknown'}</p></section>"
    ]
    for spec in storyboard:
        if isinstance(spec, _TextSlide):
            lines = "".join(
                f"<p class=\"bullet\">• {text(value)}</p>" if bullet else f"<p>{text(value)}</p>"
                for value, bullet in spec.lines
            )
            slides.append(
                f"<section class=\"slide\"><h2>{text(spec.title)}</h2>"
                f"<div class=\"body\">{lines}</div></section>"
            )
        else:
            rows = "".join(
                "<tr>" + "".join(f"<th>{text(cell)}</th>" if row_index == 0 else f"<td>{text(cell)}</td>" for cell in row) + "</tr>"
                for row_index, row in enumerate(spec.rows)
            )
            slides.append(
                f"<section class=\"slide\"><h2>{text(spec.title)}</h2>"
                f"<table>{rows}</table></section>"
            )
    return (
        "<!DOCTYPE html><html><head><meta charset=\"utf-8\"><style>"
        "@page{size:13.333333in 7.5in;margin:0}*{box-sizing:border-box}"
        "body{margin:0;background:#ddd;font-family:'Microsoft YaHei',sans-serif;color:#252525}"
        ".slide{width:13.333333in;height:7.5in;overflow:hidden;page-break-after:always;"
        "background:#f7f4ee;padding:.45in .7in}.slide:last-child{page-break-after:auto}"
        ".title-slide{display:flex;flex-direction:column;justify-content:center;"
        "align-items:center;text-align:center}.title-slide h1{font-size:30pt}"
        "h1,h2{margin:0;color:#910101}h2{font-size:25pt}.body{margin:.25in .2in;font-size:18pt;line-height:1.35}"
        "p{margin:0 0 8pt}.bullet{padding-left:.2in}table{width:100%;margin-top:.3in;border-collapse:collapse;font-size:13pt}"
        "th,td{border:1px solid #b9b1a6;padding:7pt;text-align:left;vertical-align:top}th{background:#eee8de}"
        "</style></head><body>" + "".join(slides) + "</body></html>"
    )


def _background(slide, color) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _textbox(slide, text, left, top, width, height, size, color, *, bold=False, align=None) -> None:
    from pptx.util import Inches, Pt

    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Microsoft YaHei"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    if align is not None:
        paragraph.alignment = align


def _plain(value: str) -> str:
    value = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", value)
    value = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", value)
    return re.sub(r"(`{1,2}|\*\*|__|~~|\*|_)", "", value).strip()


def _chunks(value: str, size: int) -> tuple[str, ...]:
    return tuple(value[index:index + size] for index in range(0, len(value), size) if value[index:index + size])


def _visual_line_units(value: str, *, bullet: bool) -> int:
    """Estimate wrapped 18pt CJK lines plus one paragraph-spacing unit."""

    chars_per_line = 38 if bullet else 40
    wrapped_lines = max(1, (len(value) + chars_per_line - 1) // chars_per_line)
    return wrapped_lines + 1


def _assert_table_fit(rows: tuple[tuple[str, ...], ...]) -> None:
    """Fail closed unless every cell fits the fixed editable table geometry."""

    if not rows or not rows[0]:
        raise DocumentPptxRenderError("table is empty")
    column_width_inches = 11.85 / len(rows[0])
    row_height_inches = 4.9 / len(rows)
    chars_per_line = max(1, int((column_width_inches - 0.25) / 0.18))
    visible_lines = max(1, int((row_height_inches - 0.12) / 0.24))
    capacity = chars_per_line * visible_lines
    if any(len(cell) > capacity for row in rows for cell in row):
        raise DocumentPptxRenderError("table cell exceeds visible PPTX layout budget")


def _table_cells(value: str) -> list[str]:
    return [cell.strip() for cell in value.strip().strip("|").split("|")]


def _is_table_separator(value: str) -> bool:
    cells = _table_cells(value)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _canonicalize(source: bytes) -> bytes:
    output = BytesIO()
    with ZipFile(BytesIO(source), "r") as source_zip, ZipFile(
        output, "w", compression=ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name in sorted(source_zip.namelist()):
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = (0o755 if name.endswith("/") else 0o600) << 16
            target.writestr(info, source_zip.read(name))
    return output.getvalue()
