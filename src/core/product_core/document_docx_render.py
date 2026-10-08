"""Deterministic editable DOCX rendering for governed Document Delivery."""

from __future__ import annotations

import re
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo


class DocumentDocxRenderError(ValueError):
    """Raised when the frozen Document cannot be rendered as DOCX."""


@dataclass(frozen=True, slots=True)
class DocumentDocxRenderResult:
    content: bytes
    title: str
    document_id: str
    revision: int | None


@dataclass(frozen=True, slots=True)
class DocumentDocxRenderer:
    """Render the supported Markdown subset into deterministic DOCX bytes."""

    def render(
        self,
        document: Mapping[str, object],
        markdown: str,
    ) -> DocumentDocxRenderResult:
        if not isinstance(document, Mapping):
            raise DocumentDocxRenderError("document must be a mapping")
        if not isinstance(markdown, str):
            raise DocumentDocxRenderError("markdown must be a string")
        document_id = str(document.get("id") or "")
        title = str(document.get("title") or "").strip()
        if not document_id or not title:
            raise DocumentDocxRenderError("document identity and title are required")
        revision = document.get("revision")
        if isinstance(revision, bool) or not isinstance(revision, int):
            revision = None
        try:
            content = _build_docx(
                title=title,
                document_id=document_id,
                revision=revision,
                markdown=markdown,
            )
        except DocumentDocxRenderError:
            raise
        except Exception as exc:
            raise DocumentDocxRenderError("DOCX rendering failed") from exc
        return DocumentDocxRenderResult(content, title, document_id, revision)


def _build_docx(*, title: str, document_id: str, revision: int | None, markdown: str) -> bytes:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Cm, Pt

    output = BytesIO()
    result = Document()
    section = result.sections[0]
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)

    normal = result.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    for level in range(1, 5):
        style = result.styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
    if "Code Block" not in result.styles:
        code_style = result.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "Consolas"
        code_style.font.size = Pt(9)

    properties = result.core_properties
    properties.title = title
    properties.subject = "Chriptmas OS governed Document Delivery"
    properties.identifier = document_id
    properties.version = str(revision or "")
    fixed_time = datetime(2000, 1, 1, tzinfo=UTC)
    properties.created = fixed_time
    properties.modified = fixed_time

    _append_markdown(result, markdown)
    result.save(output)
    return _canonicalize_docx(output.getvalue())


def _append_markdown(document, markdown: str) -> None:
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("```"):
            code = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            index += 1
            document.add_paragraph("\n".join(code), style="Code Block")
            continue
        heading = re.fullmatch(r"(#{1,4})\s+(.+)", stripped)
        if heading:
            document.add_heading(_plain_inline(heading.group(2)), level=len(heading.group(1)))
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and _table_separator(lines[index + 1]):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                if not _table_separator(lines[index]):
                    rows.append(_table_cells(lines[index]))
                index += 1
            width = max((len(row) for row in rows), default=0)
            if width:
                table = document.add_table(rows=len(rows), cols=width)
                table.style = "Table Grid"
                for row_index, row in enumerate(rows):
                    for cell_index, value in enumerate(row):
                        table.cell(row_index, cell_index).text = _plain_inline(value)
            continue
        unordered = re.fullmatch(r"[-+*]\s+(.+)", stripped)
        if unordered:
            document.add_paragraph(_plain_inline(unordered.group(1)), style="List Bullet")
            index += 1
            continue
        ordered = re.fullmatch(r"\d+\.\s+(.+)", stripped)
        if ordered:
            document.add_paragraph(_plain_inline(ordered.group(1)), style="List Number")
            index += 1
            continue
        if stripped.startswith(">"):
            document.add_paragraph(_plain_inline(stripped[1:].strip()), style="Quote")
            index += 1
            continue
        if stripped in {"---", "***", "___"}:
            document.add_paragraph("• • •")
            index += 1
            continue
        paragraph = [stripped]
        index += 1
        while index < len(lines) and lines[index].strip() and not _starts_block(lines, index):
            paragraph.append(lines[index].strip())
            index += 1
        document.add_paragraph(_plain_inline("\n".join(paragraph)))


def _starts_block(lines: list[str], index: int) -> bool:
    value = lines[index].strip()
    return bool(
        value.startswith(("#", "```", ">", "|"))
        or re.match(r"(?:[-+*]|\d+\.)\s+", value)
        or value in {"---", "***", "___"}
    )


def _table_separator(value: str) -> bool:
    cells = _table_cells(value)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def _table_cells(value: str) -> list[str]:
    return [cell.strip() for cell in value.strip().strip("|").split("|")]


def _plain_inline(value: str) -> str:
    value = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", value)
    value = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", value)
    value = re.sub(r"(`{1,2}|\*\*|__|~~|\*|_)", "", value)
    return value


def _canonicalize_docx(source: bytes) -> bytes:
    output = BytesIO()
    with ZipFile(BytesIO(source), "r") as archive, ZipFile(
        output, "w", compression=ZIP_DEFLATED, compresslevel=9
    ) as canonical:
        for name in sorted(archive.namelist()):
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = (0o600 if not name.endswith("/") else 0o755) << 16
            canonical.writestr(info, archive.read(name))
    return output.getvalue()
