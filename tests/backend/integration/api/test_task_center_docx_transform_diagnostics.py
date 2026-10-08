"""Source-level diagnostic reproduction for 51 local DOCX task transformations.

This isolates the production HTTP admission, SQLite Effect recovery, document
repository and task reader. Stage instrumentation sits in test-only wrappers
outside the production handler's generic failure boundary, so a failure
reports the precise stage and exception class without copying any content.
"""

from __future__ import annotations

import base64
import io
import logging
import time
from threading import Event, Thread
from collections.abc import Callable
from pathlib import Path
from types import SimpleNamespace

import pytest

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.api.container import get_container
from backend.api.job_execution_runtime import register_job_execution_handler
from backend.api.job_runtime import build_rebuild_job_repository
from backend.api.rebuild_storage_runtime import build_rebuild_object_store
from backend.api.routes import rebuild, tasks, workbench_auto_intake, workbench_original_asset
from backend.api.task_reference_projection import task_ref_for_workbench_content_transform
from backend.api.workbench_content_transform_runtime import (
    WorkbenchContentTransformDomain,
    _log_document_transform_failure,
)
from core.document_engine.sqlite_runtime import SQLiteDocumentRepository
from core.effect_log import EffectRecoveryCoordinator, build_effect_runtime
from core.product_core.source_content_read import ReadSourceTextContent
from core.product_core.source_output_memory_candidate import CreateMemoryCandidateFromSourceOutput
from core.product_core.source_structuring import StructureSourceContent


_DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_TOTAL = 51


class _SensitiveDocumentFailure(RuntimeError):
    pass


def test_document_transform_failure_log_has_only_stage_and_exception_class(tmp_path: Path, monkeypatch, caplog) -> None:
    secret = "body=private-body path=C:\\private\\source.docx credential=token-value"

    def fail_read(*_args, **_kwargs):
        raise _SensitiveDocumentFailure(secret)

    monkeypatch.setattr(ReadSourceTextContent, "execute", fail_read)
    domain = WorkbenchContentTransformDomain(tmp_path, object(), "default", object())
    item = {"source_id": "source-document-log", "project_id": "default", "source_title": "local-docx"}
    with caplog.at_level(logging.WARNING, logger="backend.api.workbench_content_transform_runtime"):
        with pytest.raises(_SensitiveDocumentFailure, match="private-body"):
            domain._execute_document(item, execution_ref="facts:effect/test", checkpoint=lambda: None)

    messages = [record.getMessage() for record in caplog.records if record.name == "backend.api.workbench_content_transform_runtime"]
    assert messages == ["workbench_document_transform_failed stage=read exception_type=_SensitiveDocumentFailure"]
    assert all(record.exc_info is None for record in caplog.records)
    assert secret not in caplog.text
    assert "private-body" not in caplog.text
    assert "source.docx" not in caplog.text
    assert "token-value" not in caplog.text


def test_document_transform_failure_log_bounds_dynamic_exception_class(caplog) -> None:
    long_exception = type("X" * 160, (RuntimeError,), {})
    with caplog.at_level(logging.WARNING, logger="backend.api.workbench_content_transform_runtime"):
        _log_document_transform_failure("read", long_exception("body=private-body credential=token-value"))

    messages = [record.getMessage() for record in caplog.records if record.name == "backend.api.workbench_content_transform_runtime"]
    assert messages == [f"workbench_document_transform_failed stage=read exception_type={'X' * 80}"]
    assert all(record.exc_info is None for record in caplog.records)
    assert "private-body" not in caplog.text
    assert "token-value" not in caplog.text


class _NoRemoteServices:
    def __getattr__(self, name: str):
        raise AssertionError(f"local DOCX diagnostics must not access services or secrets: {name}")


def _application(root: Path) -> FastAPI:
    application = FastAPI()
    for router in (workbench_original_asset.router, workbench_auto_intake.router, tasks.router, rebuild.router):
        application.include_router(router)
    container = SimpleNamespace(root_dir=root, settings_service=_NoRemoteServices(), secret_store=_NoRemoteServices())
    application.dependency_overrides[get_container] = lambda: container
    runtime = build_effect_runtime(root / ".rebuild-data" / "jobs.sqlite3", owner_id="docx-51-diagnostics")
    application.state.effect_runtime = runtime
    register_job_execution_handler(application, root, runtime)
    return application


def _docx(index: int) -> bytes:
    document = Document()
    document.add_heading(f"DOCX 任务 {index:02d}", level=1)
    document.add_paragraph(f"第 {index:02d} 份本地受控 DOCX，用于验证任务分页和转换交付。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "编号"
    table.cell(0, 1).text = str(index)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _submit(client: TestClient, index: int) -> str:
    content = _docx(index)
    filename = f"pagination-{index:03d}.docx"
    uploaded = client.post("/api/rebuild/workbench/original-asset", json={
        "display_name": filename,
        "media_type": _DOCX_TYPE,
        "size_bytes": len(content),
        "content_base64": base64.b64encode(content).decode("ascii"),
        "source_kind": "file",
    })
    assert uploaded.status_code == 201, uploaded.json()
    admission = client.post("/api/rebuild/workbench/auto-intake", json={
        "content": "",
        "media_type": _DOCX_TYPE,
        "file_name": filename,
        "title": filename,
        "original_asset_ref": uploaded.json()["asset_ref"],
        "add_to_knowledge_base": True,
    })
    assert admission.status_code == 202, admission.json()
    return admission.json()["job_id"]


def _install_stage_recorder(monkeypatch) -> tuple[dict[str, str], list[dict[str, str]]]:
    stages: dict[str, str] = {}
    failures: list[dict[str, str]] = []

    def wrap_method(target, attribute: str, stage: str) -> None:
        original = getattr(target, attribute)

        def wrapped(instance, *args, **kwargs):
            source_id = kwargs.get("source_id")
            if not isinstance(source_id, str):
                source_id = ""
            subject = args[0] if args else None
            if not source_id and isinstance(subject, str):
                source_id = subject
            if not source_id:
                source_refs = getattr(subject, "source_refs", ())
                if source_refs and isinstance(source_refs[0], dict):
                    ref_source_id = source_refs[0].get("source_id")
                    source_id = ref_source_id if isinstance(ref_source_id, str) else ""
            stages[source_id] = stage
            try:
                return original(instance, *args, **kwargs)
            except Exception as error:
                failures.append({
                    "source_id": source_id,
                    "stage": stage,
                    "exception_type": type(error).__name__,
                })
                raise

        monkeypatch.setattr(target, attribute, wrapped)

    wrap_method(ReadSourceTextContent, "execute", "read_source_text")
    wrap_method(StructureSourceContent, "execute", "structure_source_content")
    wrap_method(CreateMemoryCandidateFromSourceOutput, "execute_from_content_read", "create_memory_candidate")
    wrap_method(SQLiteDocumentRepository, "create_or_replay_generated", "create_document")

    original_execute = WorkbenchContentTransformDomain.execute_item

    def execute_item(domain, effect, item, checkpoint):
        source_id = str(item["source_id"])
        stages[source_id] = "dispatch_transform"
        try:
            return original_execute(domain, effect, item, checkpoint)
        except Exception as error:
            failures.append({
                "source_id": source_id,
                "stage": stages.get(source_id, "dispatch_transform"),
                "exception_type": type(error).__name__,
            })
            raise

    monkeypatch.setattr(WorkbenchContentTransformDomain, "execute_item", execute_item)
    original_verify = WorkbenchContentTransformDomain.verify_output

    def verify_output(domain, effect, item, output):
        source_id = str(item["source_id"])
        stages[source_id] = "verify_output"
        try:
            return original_verify(domain, effect, item, output)
        except Exception as error:
            failures.append({
                "source_id": source_id,
                "stage": "verify_output",
                "exception_type": type(error).__name__,
            })
            raise

    monkeypatch.setattr(WorkbenchContentTransformDomain, "verify_output", verify_output)
    return stages, failures


def test_document_auto_intake_candidate_uses_bounded_source_preview_and_stays_pending(
    tmp_path: Path,
) -> None:
    application = _application(tmp_path)
    store, _settings = build_rebuild_object_store(tmp_path)
    repository = build_rebuild_job_repository(tmp_path, store)
    with TestClient(application) as client:
        job_id = _submit(client, 1)
    coordinator = EffectRecoveryCoordinator(application.state.effect_runtime)
    for offset in (1, 61, 121):
        coordinator.recover_once(now=int(time.time()) + offset)
        job = repository.get(job_id)
        if job is not None and job["status"] in {"completed", "failed", "waiting_user", "cancelled"}:
            break
    else:
        raise AssertionError("controlled DOCX transform did not reach a terminal state")

    assert job is not None and job["status"] == "completed"
    candidates = [
        item for item in store.list("memory_candidates")
        if item.get("candidate_type") == "document_takeaway"
    ]
    assert len(candidates) == 1
    candidate = candidates[0]
    assert candidate["project_id"] == "default"
    assert candidate["status"] == "pending_review"
    assert "DOCX 任务 01" in candidate["proposed_content"]
    assert "第 01 份本地受控 DOCX" in candidate["proposed_content"]
    assert str(candidate["source_refs"][0]["source_id"]).startswith("source-file-")
    assert candidate["review"]["requires_user_confirmation"] is True
    assert candidate["review"]["auto_promote_allowed"] is False


def test_fifty_one_docx_transforms_report_precise_pre_generic_failure_stage(tmp_path: Path, monkeypatch) -> None:
    stages, failures = _install_stage_recorder(monkeypatch)
    application = _application(tmp_path)
    store, _settings = build_rebuild_object_store(tmp_path)
    repository = build_rebuild_job_repository(tmp_path, store)
    coordinator = EffectRecoveryCoordinator(application.state.effect_runtime)
    terminal = {"completed", "failed", "waiting_user", "cancelled"}
    stop = Event()
    recovery_errors: list[str] = []

    def recover_while_admitting() -> None:
        while not stop.is_set():
            try:
                coordinator.recover_once(now=int(time.time()) + 1)
            except Exception as error:
                recovery_errors.append(type(error).__name__)
            stop.wait(0.2)

    recovery = Thread(target=recover_while_admitting, name="docx-51-core-recovery", daemon=True)
    recovery.start()
    job_ids: list[str] = []
    try:
        with TestClient(application) as client:
            for index in range(1, _TOTAL + 1):
                job_ids.append(_submit(client, index))
                # Match the packaged UI's paced import while a separate Core
                # recovery tick may execute earlier admissions concurrently.
                time.sleep(0.3)

            deadline = time.monotonic() + 30
            while time.monotonic() < deadline:
                jobs = [repository.get(job_id) for job_id in job_ids]
                if all(job is not None and job["status"] in terminal for job in jobs):
                    break
                time.sleep(0.2)
            else:
                raise AssertionError("51 local DOCX jobs did not reach terminal states")
    finally:
        stop.set()
        recovery.join(timeout=5)

    jobs = [repository.get(job_id) for job_id in job_ids]
    failed = [job for job in jobs if job is None or job["status"] != "completed"]
    assert not recovery_errors
    assert not failed, {"failed_jobs": failed, "stage_failures": failures}
    assert len(stages) == _TOTAL
    assert not failures

    with TestClient(application) as client:
        first = client.get("/api/rebuild/tasks", params={"project_id": "default", "limit": 50})
        assert first.status_code == 200, first.json()
        assert len(first.json()["items"]) == 50
        assert first.json()["next_cursor"]
        second = client.get("/api/rebuild/tasks", params={
            "project_id": "default", "limit": 50, "cursor": first.json()["next_cursor"],
        })
        assert second.status_code == 200, second.json()
        assert len(second.json()["items"]) == 1
        listed_refs = {item["task_ref"] for item in first.json()["items"] + second.json()["items"]}
        assert listed_refs == {
            task_ref_for_workbench_content_transform(project_id="default", job_id=job_id)
            for job_id in job_ids
        }
        assert {item["status"] for item in first.json()["items"] + second.json()["items"]} == {"delivered"}

