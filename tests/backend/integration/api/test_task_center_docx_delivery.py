"""Controlled DOCX delivery through production admission, recovery and HTTP readers.

The transport is in-process and the settings/secret ports deliberately cannot
contact a provider. The DOCX parser, file store, SQLite execution, document
repository, task queries and edit-conflict responses are real implementations.
This is isolated vertical evidence, not an installed-candidate acceptance.
"""

from __future__ import annotations

import base64
import io
import time
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.api.container import get_container
from backend.api.job_execution_runtime import register_job_execution_handler
from backend.api.job_runtime import build_rebuild_job_repository
from backend.api.rebuild_storage_runtime import build_rebuild_object_store
from backend.api.routes import rebuild, tasks, workbench_auto_intake, workbench_original_asset
from backend.api.routes.product import repositories as product_repositories
from backend.api.task_reference_projection import task_ref_for_workbench_content_transform
from backend.memory_app.app import create_app as add_workspace
from core.effect_log import EffectRecoveryCoordinator, build_effect_runtime
from core.product_core.workbench_content_transform_execution import read_workbench_transform_receipt


_DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_BODY = "任务化验收文档。第一项记录原始来源。第二项修改后仍可追溯。"
_TABLE = "任务验收表格正文"


class _NoRemoteServices:
    def __getattr__(self, name):
        raise AssertionError(f"local DOCX delivery must not access services or secrets: {name}")


def _application(root: Path, *, owner: str) -> FastAPI:
    application = FastAPI()
    for router in (workbench_original_asset.router, workbench_auto_intake.router, tasks.router, rebuild.router):
        application.include_router(router)
    container = SimpleNamespace(
        root_dir=root, settings_service=_NoRemoteServices(), secret_store=_NoRemoteServices(),
    )
    application.dependency_overrides[get_container] = lambda: container
    runtime = build_effect_runtime(root / ".rebuild-data" / "jobs.sqlite3", owner_id=owner)
    application.state.effect_runtime = runtime
    register_job_execution_handler(application, root, runtime)
    return add_workspace(runtime_root=root, legacy_app=application)


def _controlled_docx() -> bytes:
    document = Document()
    document.add_heading("受控任务验收", level=1)
    document.add_paragraph(_BODY)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "来源"
    table.cell(0, 1).text = _TABLE
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _submit(client: TestClient) -> tuple[dict, dict]:
    content = _controlled_docx()
    uploaded = client.post("/api/rebuild/workbench/original-asset", json={
        "display_name": "受控任务验收.docx", "media_type": _DOCX_TYPE,
        "size_bytes": len(content), "content_base64": base64.b64encode(content).decode("ascii"),
        "source_kind": "file",
    })
    assert uploaded.status_code == 201, uploaded.json()
    body = {
        "content": "", "media_type": _DOCX_TYPE, "file_name": "受控任务验收.docx",
        "title": "受控任务验收", "original_asset_ref": uploaded.json()["asset_ref"],
        "add_to_knowledge_base": True,
    }
    response = client.post("/api/rebuild/workbench/auto-intake", json=body)
    assert response.status_code == 202, response.json()
    return response.json(), body


def _recover(application: FastAPI, root: Path, job_id: str) -> dict:
    store, _settings = build_rebuild_object_store(root)
    now = int(time.time()) + 1
    for attempt in range(3):
        EffectRecoveryCoordinator(application.state.effect_runtime).recover_once(now=now + attempt * 60)
        job = build_rebuild_job_repository(root, store).get(job_id)
        assert job is not None
        if job["status"] in {"completed", "failed", "waiting_user", "cancelled"}:
            return job
    raise AssertionError("controlled DOCX did not reach a terminal state after Core recovery")


def test_plain_text_auto_intake_does_not_claim_a_transform_task_reference(tmp_path: Path):
    application = _application(tmp_path, owner="plain-text-admission")
    with TestClient(application) as client:
        response = client.post("/api/rebuild/workbench/auto-intake", json={
            "content": "这是一条普通文本记录，不应伪装成资料转换任务。",
            "add_to_knowledge_base": True,
        })
    assert response.status_code in {201, 202}, response.json()
    assert response.json()["job_id"].startswith("job-intake-")
    assert "task_ref" not in response.json()
    assert response.json()["project_id"] == "default"


def test_docx_closing_client_does_not_stop_delivery_and_restarting_preserves_edit_conflict(tmp_path: Path):
    first = _application(tmp_path, owner="docx-intake-client")
    with TestClient(first) as client:
        admission, body = _submit(client)
        job_id = admission["job_id"]
        source_id = admission["items"][0]["source_id"]
        expected_ref = task_ref_for_workbench_content_transform(
            project_id="default", job_id=job_id,
        )
        assert admission["project_id"] == "default"
        assert admission["task_ref"] == expected_ref
        store, _settings = build_rebuild_object_store(tmp_path)
        assert not store.list("source_content_reads")
        replay = client.post("/api/rebuild/workbench/auto-intake", json=body)
        assert replay.status_code in {200, 202}, replay.json()
        assert replay.json()["job_id"] == job_id
        assert replay.json()["project_id"] == "default"
        assert replay.json()["task_ref"] == expected_ref

    # No client or page drives the continuation. A newly composed Core discovers
    # durable admission using the same isolated disk state.
    restarted = _application(tmp_path, owner="docx-restarted-core")
    job = _recover(restarted, tmp_path, job_id)
    assert job["status"] == "completed", job
    receipt = read_workbench_transform_receipt(tmp_path / ".rebuild-data" / "jobs.sqlite3", job_id)
    assert receipt is not None
    output = receipt["outputs"][0]
    document_id = output["document_id"]

    with TestClient(restarted) as client:
        listed = client.get("/api/rebuild/tasks", params={"project_id": "default"})
        assert listed.status_code == 200, listed.json()
        assert any(item["task_ref"] == expected_ref for item in listed.json()["items"]), {
            "listed": listed.json(),
            "source_scope": {key: store.read("sources", source_id).get(key) for key in ("id", "project_id", "type")},
            "job_scope": {key: job.get(key) for key in ("id", "project_id", "job_type", "execution_version", "status")},
        }
        task = next(item for item in listed.json()["items"] if item["task_ref"] == expected_ref)
        task_ref = task["task_ref"]
        detail_url = f"/api/rebuild/tasks/{task_ref}?project_id=default"
        detail = client.get(detail_url)
        assert detail.status_code == 200, detail.json()
        assert detail.json()["status"] == "delivered", {
            "detail": detail.json(), "output": output,
            "published": job.get("published_outputs"),
            "document": client.get(f"/api/rebuild/documents/{document_id}").json(),
        }
        assert any(item["artifact_id"] == document_id for item in detail.json()["detail"]["outputs"])
        assert detail.json()["detail"]["outputs"][0]["kind"] == "review"
        assert detail.json()["detail"]["review_items"][0]["source_id"] == source_id

        opened = client.get(f"/api/rebuild/documents/{document_id}")
        assert opened.status_code == 404
        premature_candidate = client.post("/api/rebuild/memory/candidates/review", json={
            "candidate_id": output["candidate_id"], "action": "confirm",
        })
        assert premature_candidate.status_code == 409, premature_candidate.json()
        review = client.get("/api/workspace/v1/legacy-reviews", params={"project_id": "default"})
        assert review.status_code == 200, review.text
        assert review.json()["items"][0]["source_id"] == source_id
        assert review.json()["items"][0]["status"] == "ready"
        assert _BODY in review.json()["items"][0]["source_text"]
        generated = product_repositories._document_repository(tmp_path, store, build_rebuild_object_store(tmp_path)[1]).markdown(document_id)
        assert _BODY in generated and _TABLE in generated
        reviewed = generated + "\n\n用户核对：保留原始来源。"
        saved_draft = client.put(f"/api/workspace/v1/legacy-reviews/{source_id}/draft", json={
            "project_id": "default", "markdown": reviewed,
        })
        assert saved_draft.status_code == 200, saved_draft.text
        confirmed = client.post(f"/api/workspace/v1/legacy-reviews/{source_id}/confirm", json={
            "project_id": "default",
        })
        assert confirmed.status_code == 200, confirmed.text
        assert confirmed.json()["document_id"] == document_id
        assert confirmed.json()["document_revision"] == 2
        opened = client.get(f"/api/rebuild/documents/{document_id}")
        assert opened.status_code == 200, opened.json()
        original_document = opened.json()
        assert original_document["source_task_ref"] == task_ref
        assert _BODY in original_document["markdown"]
        assert _TABLE in original_document["markdown"]
        assert original_document["markdown"] == reviewed
        assert any(ref["source_id"] == source_id for ref in original_document["source_refs"])
        revision = original_document["revision"]
        changed = original_document["markdown"] + "\n\n用户核对后补充：保留原始来源。"
        saved = client.put(f"/api/rebuild/documents/{document_id}", json={
            "expected_revision": revision, "markdown": changed,
        })
        assert saved.status_code == 200, saved.json()
        assert saved.json()["revision"] == revision + 1
        assert saved.json()["source_task_ref"] == task_ref
        rediscovered = client.get(detail_url)
        assert rediscovered.status_code == 200, rediscovered.json()
        assert rediscovered.json()["detail"]["outputs"][0]["kind"] == "document"
        conflict = client.put(f"/api/rebuild/documents/{document_id}", json={
            "expected_revision": revision, "markdown": "另一个客户端的旧草稿",
        })
        assert conflict.status_code == 409, conflict.json()
        assert conflict.json()["current_document"]["markdown"] == changed
        assert conflict.json()["current_document"]["source_task_ref"] == task_ref

    final = _application(tmp_path, owner="docx-after-edit-restart")
    with TestClient(final) as client:
        reopened = client.get(f"/api/rebuild/documents/{document_id}")
        assert reopened.json()["markdown"] == changed
        assert reopened.json()["revision"] == revision + 1
        assert reopened.json()["source_task_ref"] == task_ref
        rediscovered = client.get(detail_url)
        assert rediscovered.status_code == 200, rediscovered.json()
        assert rediscovered.json()["status"] == "delivered", rediscovered.json()
    store, _settings = build_rebuild_object_store(tmp_path)
    candidate = store.read("memory_candidates", output["candidate_id"])
    assert candidate["status"] == "pending_review"
    assert len(store.list("sources")) == 1
