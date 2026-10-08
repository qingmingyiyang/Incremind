from __future__ import annotations

import time
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from core.effect_log import EffectReaper

from core.document_engine import DocumentDraft, ObjectStoreDocumentRepository
from core.product_core.document_html_render import DocumentHtmlRenderer
from core.product_core.document_docx_render import DocumentDocxRenderer
from core.product_core.document_pptx_render import (
    DocumentPptxRenderError,
    DocumentPptxRenderer,
)
from core.product_core.document_delivery import (
    DocumentDeliveryConflict,
    DocumentDeliveryError,
    DocumentDeliveryService,
    DocumentDeliveryStyleRegistry,
)


def _recover_with_core(service: DocumentDeliveryService) -> dict[str, object]:
    EffectReaper(service._effects).recover_expired(
        now=int(time.time()) + 301,
        verifiers={"document_delivery": lambda effect: service.verify_effect(effect.operation_id)},
    )
    planned = service._effects.planned_for_kinds(("document_delivery",), limit=64)
    recovered: list[str] = []
    failed: list[dict[str, str]] = []
    for effect in planned:
        try:
            service._runner.execute_planned(
                effect.operation_id,
                service.handle_effect,
                now=int(time.time()) + 301,
                receipt_kind="document-delivery-receipt",
            )
        except Exception as error:
            failed.append({"delivery_id": effect.operation_id, "error": str(error)})
        else:
            recovered.append(effect.operation_id)
    return {
        "attempted": len(planned),
        "recovered": recovered,
        "failed": failed,
    }
from core.storage_provider import JsonObjectStore, SQLiteStructuredRecordStore


def _parts(tmp_path: Path):
    object_store = JsonObjectStore(tmp_path / ".rebuild-data", legacy_root=tmp_path / "library")
    documents = ObjectStoreDocumentRepository(object_store, now="2026-08-26T00:00:00Z")
    document = documents.create(DocumentDraft(
        document_type="media_analysis",
        title="Media result",
        markdown="# Media result\n\nExact body.",
        source_refs=({
            "source_id": "source-one",
            "locator": "crp://default/sources/source-one",
        },),
        project_id="project-one",
    ))
    records = SQLiteStructuredRecordStore(tmp_path / ".rebuild-data" / "jobs.sqlite3")
    records.list("document_deliveries")
    service = DocumentDeliveryService(records, documents, tmp_path)
    return service, records, documents, document


def test_delivery_freezes_revision_style_and_publishes_markdown_html(tmp_path: Path) -> None:
    service, records, _documents, document = _parts(tmp_path)

    result = service.create_or_resume(
        document_id=str(document["id"]),
        expected_document_revision=1,
        formats=["html", "markdown"],
    )

    assert result["status"] == "completed" and result["replayed"] is False
    assert result["document_revision"] == 1
    assert result["formats"] == ["markdown", "html"]
    assert result["style_snapshot"] == {
        "profile_id": "builtin.porcelain-document",
        "revision": 1,
        "renderer_id": "rebuild.document-html",
        "renderer_revision": 1,
        "delivery_key": "porcelain-html-v1",
    }
    assert {item["format"] for item in result["artifacts"]} == {"markdown", "html"}
    assert all(item["verified"] is True for item in result["artifacts"])
    assert "file_path" not in str(result) and "Exact body" not in str(result)
    stored = records.read("document_deliveries", str(result["delivery_id"]))
    assert stored is not None
    assert "markdown" not in stored.payload["document_snapshot"]
    receipt = records.read("document_delivery_receipts", str(result["delivery_id"]))
    assert receipt is not None and receipt.payload["status"] == "completed"
    assert result["receipt_ref"].startswith("crp://default/document-delivery-receipts/")

    md_name, markdown = service.artifact(str(result["delivery_id"]), "markdown")
    html_name, html = service.artifact(str(result["delivery_id"]), "html")
    assert md_name.endswith(".md") and markdown == b"# Media result\n\nExact body."
    assert html_name.endswith(".html") and b"<!DOCTYPE html>" in html


def test_docx_renderer_is_deterministic_editable_and_structured() -> None:
    from docx import Document

    renderer = DocumentDocxRenderer()
    markdown = """# Report

Paragraph with **bold source**.

- First
- Second

| Name | Value |
| --- | --- |
| A | 1 |

```text
exact code
```
"""
    document = {"id": "document-report", "title": "Report", "revision": 3}

    first = renderer.render(document, markdown).content
    second = renderer.render(document, markdown).content
    parsed = Document(BytesIO(first))

    assert first == second
    assert first.startswith(b"PK")
    assert any(paragraph.text == "Report" for paragraph in parsed.paragraphs)
    assert any(paragraph.text == "First" for paragraph in parsed.paragraphs)
    assert parsed.tables[0].cell(1, 0).text == "A"
    with ZipFile(BytesIO(first)) as archive:
        assert all(info.date_time == (1980, 1, 1, 0, 0, 0) for info in archive.infolist())


def test_delivery_publishes_and_replays_verified_docx(tmp_path: Path) -> None:
    from docx import Document

    service, records, _documents, document = _parts(tmp_path)
    result = service.create_or_resume(
        document_id=str(document["id"]),
        expected_document_revision=1,
        formats=["docx", "markdown", "html"],
    )

    assert result["formats"] == ["markdown", "html", "docx"]
    assert result["style_snapshot"]["docx_renderer_id"] == "rebuild.document-docx"
    assert result["style_snapshot"]["docx_renderer_revision"] == 1
    file_name, content = service.artifact(str(result["delivery_id"]), "docx")
    assert file_name.endswith(".docx")
    assert "Exact body." in "\n".join(
        paragraph.text for paragraph in Document(BytesIO(content)).paragraphs
    )
    replay = service.create_or_resume(
        document_id=str(document["id"]),
        expected_document_revision=1,
        formats=["markdown", "html", "docx"],
    )
    assert replay["delivery_id"] == result["delivery_id"] and replay["replayed"] is True
    stored = records.read("document_deliveries", str(result["delivery_id"]))
    assert stored is not None and "markdown" not in stored.payload


def test_pptx_renderer_is_deterministic_editable_and_structured() -> None:
    from pptx import Presentation

    renderer = DocumentPptxRenderer()
    markdown = """# Overview

- First point
- Second point

## Evidence

| Name | Value |
| --- | --- |
| A | 1 |
"""
    first = renderer.render({"id": "document-report", "title": "Report", "revision": 3}, markdown)
    second = renderer.render({"id": "document-report", "title": "Report", "revision": 3}, markdown)
    parsed = Presentation(BytesIO(first.content))

    assert first.content == second.content
    assert first.slide_html == second.slide_html
    assert first.slide_html.startswith("<!DOCTYPE html>")
    assert "<script" not in first.slide_html
    assert first.slide_count == len(parsed.slides) == 3
    editable_text = "\n".join(
        shape.text
        for slide in parsed.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    assert "Report" in editable_text and "First point" in editable_text
    assert any(shape.has_table for slide in parsed.slides for shape in slide.shapes)
    with ZipFile(BytesIO(first.content)) as archive:
        assert all(info.date_time == (1980, 1, 1, 0, 0, 0) for info in archive.infolist())
        assert "ppt/presentation.xml" in archive.namelist()


def test_pptx_renderer_preserves_long_text_and_repeats_split_table_header() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    long_text = "证据" * 420
    table_rows = "\n".join(f"| row-{index} | value-{index} |" for index in range(15))
    rendered = DocumentPptxRenderer(max_lines_per_slide=4).render(
        {"id": "document-long", "title": "Long report", "revision": 1},
        f"# Long section\n\n{long_text}\n\n## Table\n\n| Name | Value |\n| --- | --- |\n{table_rows}",
    )
    parsed = Presentation(BytesIO(rendered.content))
    text = "".join(
        shape.text
        for slide in parsed.slides
        for shape in slide.shapes
        if hasattr(shape, "text") and shape.top >= Inches(1.4)
    )
    assert long_text in text.replace("\n", "")
    tables = [shape.table for slide in parsed.slides for shape in slide.shapes if shape.has_table]
    assert len(tables) == 2
    assert all(table.cell(0, 0).text == "Name" for table in tables)


def test_pptx_renderer_paginates_maximum_cjk_text_by_visible_line_budget() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    long_text = "内容" * 650
    rendered = DocumentPptxRenderer().render(
        {"id": "document-visual-budget", "title": "Visual budget", "revision": 1},
        f"# Long\n\n{long_text}",
    )
    parsed = Presentation(BytesIO(rendered.content))
    body_slides = list(parsed.slides)[1:]
    assert len(body_slides) == 3
    recovered = "".join(
        shape.text.replace("\n", "")
        for slide in body_slides
        for shape in slide.shapes
        if hasattr(shape, "text") and shape.top >= Inches(1.4)
    )
    assert recovered == long_text
    assert all(
        len(shape.text) <= 480
        for slide in body_slides
        for shape in slide.shapes
        if hasattr(shape, "text") and shape.top >= Inches(1.4)
    )


def test_pptx_renderer_rejects_table_cell_that_cannot_fit_visible_geometry() -> None:
    rows = "\n".join(
        f"| row-{index} | {'很长的证据内容' * 45 if index == 10 else 'short'} |"
        for index in range(11)
    )
    with pytest.raises(DocumentPptxRenderError, match="visible PPTX layout budget"):
        DocumentPptxRenderer().render(
            {"id": "document-table-overflow", "title": "Table overflow", "revision": 1},
            f"# Table\n\n| Name | Value |\n| --- | --- |\n{rows}",
        )


def test_delivery_publishes_and_replays_verified_pptx(tmp_path: Path) -> None:
    from pptx import Presentation

    service, records, _documents, document = _parts(tmp_path)
    result = service.create_or_resume(
        document_id=str(document["id"]),
        expected_document_revision=1,
        formats=["pptx", "markdown", "html", "docx"],
    )

    assert result["formats"] == ["markdown", "html", "docx", "pptx"]
    assert result["style_snapshot"]["pptx_renderer_id"] == "rebuild.document-pptx"
    assert result["style_snapshot"]["pptx_renderer_revision"] == 2
    file_name, content = service.artifact(str(result["delivery_id"]), "pptx")
    assert file_name.endswith(".pptx")
    parsed = Presentation(BytesIO(content))
    assert len(parsed.slides) >= 1
    assert any("Media result" in shape.text for slide in parsed.slides for shape in slide.shapes if hasattr(shape, "text"))
    replay = service.create_or_resume(
        document_id=str(document["id"]),
        expected_document_revision=1,
        formats=["markdown", "html", "docx", "pptx"],
    )
    assert replay["delivery_id"] == result["delivery_id"] and replay["replayed"] is True
    stored = records.read("document_deliveries", str(result["delivery_id"]))
    assert stored is not None and "Exact body" not in str(stored.payload)
    assert stored.payload["pptx_layout_plan"]["schema_version"] == "1.0.0"
    assert stored.payload["pptx_layout_plan"]["relative_path"].endswith(
        "/layout-plan.json"
    )
    receipt = records.read(
        "document_delivery_receipts", str(result["delivery_id"])
    )
    assert receipt is not None
    assert receipt.payload["pptx_layout_plan"] == stored.payload["pptx_layout_plan"]
    slide_html = service.slide_html(str(result["delivery_id"]))
    assert slide_html.startswith("<!DOCTYPE html>") and "Media result" in slide_html
    assert "slide_html" not in str(stored.payload)


def test_frozen_pptx_layout_plan_tamper_fails_closed_without_changing_pptx(
    tmp_path: Path,
) -> None:
    service, records, _documents, document = _parts(tmp_path)
    result = service.create_or_resume(
        document_id=str(document["id"]),
        expected_document_revision=1,
        formats=["pptx"],
    )
    delivery_id = str(result["delivery_id"])
    _file_name, original_pptx = service.artifact(delivery_id, "pptx")
    stored = records.read("document_deliveries", delivery_id)
    assert stored is not None
    descriptor = stored.payload["pptx_layout_plan"]
    plan_path = (
        tmp_path / "exports" / "document-delivery" / descriptor["relative_path"]
    )
    plan_path.write_text('{"schema_version":"1.0.0"}', encoding="utf-8")

    with pytest.raises(DocumentDeliveryConflict, match="layout plan"):
        service.slide_html(delivery_id)

    assert original_pptx.startswith(b"PK")


def test_pptx_renderer_upgrade_keeps_r1_recovery_and_creates_isolated_r2(
    tmp_path: Path, monkeypatch
) -> None:
    import core.product_core.document_delivery as delivery_module

    _service, records, documents, document = _parts(tmp_path)
    style = {
        "profile_id": "builtin.porcelain-document", "revision": 1,
        "renderer_id": "rebuild.document-html", "renderer_revision": 1,
        "delivery_key": "porcelain-html-v1",
    }
    style_key = ("builtin.porcelain-document", 1, "rebuild.document-html", 1)

    class MarkerPptxRenderer:
        def __init__(self, marker: bytes) -> None:
            self.marker = marker

        def plan(self, document, markdown):
            return DocumentPptxRenderer().plan(document, markdown)

        def render(self, _document, _markdown, *, layout_plan=None):
            return SimpleNamespace(content=self.marker)

    r1_key = (*style_key, "rebuild.document-pptx", 1)
    r2_key = (*style_key, "rebuild.document-pptx", 2)

    def service_for(revision: int):
        return DocumentDeliveryService(
            records,
            documents,
            tmp_path,
            style_registry=DocumentDeliveryStyleRegistry(
                current_snapshot=style,
                renderers={style_key: DocumentHtmlRenderer()},
                delivery_keys={style_key: "porcelain-html-v1"},
                current_pptx_renderer=("rebuild.document-pptx", revision),
                pptx_renderers={
                    r1_key: MarkerPptxRenderer(b"pptx-r1"),
                    r2_key: MarkerPptxRenderer(b"pptx-r2"),
                },
                pptx_delivery_keys={r1_key: "pptx-v1", r2_key: "pptx-v2"},
            ),
        )

    real_publish = delivery_module._publish_exact

    def crash_after_publish(path: Path, data: bytes) -> None:
        real_publish(path, data)
        raise BaseException("simulated process exit")

    monkeypatch.setattr(delivery_module, "_publish_exact", crash_after_publish)
    with pytest.raises(BaseException, match="simulated process exit"):
        service_for(1).create_or_resume(
            document_id=str(document["id"]),
            expected_document_revision=1,
            formats=["pptx"],
        )
    monkeypatch.setattr(delivery_module, "_publish_exact", real_publish)

    restarted = service_for(2)
    recovery = _recover_with_core(restarted)
    r1_delivery_id = str(recovery["recovered"][0])
    r2 = restarted.create_or_resume(
        document_id=str(document["id"]),
        expected_document_revision=1,
        formats=["pptx"],
    )
    assert restarted.artifact(r1_delivery_id, "pptx")[1] == b"pptx-r1"
    assert restarted.artifact(str(r2["delivery_id"]), "pptx")[1] == b"pptx-r2"
    assert r1_delivery_id != r2["delivery_id"]


def test_docx_renderer_upgrade_keeps_r1_recovery_and_creates_isolated_r2(
    tmp_path: Path, monkeypatch
) -> None:
    import core.product_core.document_delivery as delivery_module

    _service, records, documents, document = _parts(tmp_path)
    style = {
        "profile_id": "builtin.porcelain-document", "revision": 1,
        "renderer_id": "rebuild.document-html", "renderer_revision": 1,
        "delivery_key": "porcelain-html-v1",
    }
    style_key = (
        "builtin.porcelain-document", 1, "rebuild.document-html", 1,
    )

    class MarkerDocxRenderer:
        def __init__(self, marker: bytes) -> None:
            self.marker = marker

        def render(self, _document, _markdown):
            return SimpleNamespace(content=self.marker)

    r1_key = (*style_key, "rebuild.document-docx", 1)
    r2_key = (*style_key, "rebuild.document-docx", 2)

    def service_for(revision: int):
        return DocumentDeliveryService(
            records,
            documents,
            tmp_path,
            style_registry=DocumentDeliveryStyleRegistry(
                current_snapshot=style,
                renderers={style_key: DocumentHtmlRenderer()},
                delivery_keys={style_key: "porcelain-html-v1"},
                current_docx_renderer=("rebuild.document-docx", revision),
                docx_renderers={
                    r1_key: MarkerDocxRenderer(b"docx-r1"),
                    r2_key: MarkerDocxRenderer(b"docx-r2"),
                },
                docx_delivery_keys={r1_key: "docx-v1", r2_key: "docx-v2"},
            ),
        )

    real_publish = delivery_module._publish_exact

    def crash_after_publish(path: Path, data: bytes) -> None:
        real_publish(path, data)
        raise BaseException("simulated process exit")

    monkeypatch.setattr(delivery_module, "_publish_exact", crash_after_publish)
    with pytest.raises(BaseException, match="simulated process exit"):
        service_for(1).create_or_resume(
            document_id=str(document["id"]),
            expected_document_revision=1,
            formats=["docx"],
        )
    monkeypatch.setattr(delivery_module, "_publish_exact", real_publish)

    restarted = service_for(2)
    recovery = _recover_with_core(restarted)
    r1_delivery_id = str(recovery["recovered"][0])
    r2 = restarted.create_or_resume(
        document_id=str(document["id"]),
        expected_document_revision=1,
        formats=["docx"],
    )

    assert r1_delivery_id != r2["delivery_id"]
    assert restarted.artifact(r1_delivery_id, "docx")[1] == b"docx-r1"
    assert restarted.artifact(str(r2["delivery_id"]), "docx")[1] == b"docx-r2"
    r1_record = records.read("document_deliveries", r1_delivery_id)
    r2_record = records.read("document_deliveries", str(r2["delivery_id"]))
    assert r1_record is not None and r2_record is not None
    assert r1_record.payload["style_snapshot"]["docx_delivery_key"] == "docx-v1"
    assert r2_record.payload["style_snapshot"]["docx_delivery_key"] == "docx-v2"
    assert (
        records.read("document_delivery_receipts", r1_delivery_id).payload["artifacts"][0]["relative_path"]
        != records.read("document_delivery_receipts", str(r2["delivery_id"])).payload["artifacts"][0]["relative_path"]
    )


def test_delivery_rejects_stale_revision_before_file_effect(tmp_path: Path) -> None:
    service, _records, documents, document = _parts(tmp_path)
    documents.save_user_edit(
        str(document["id"]), markdown="new", expected_revision=1, title="new"
    )

    with pytest.raises(DocumentDeliveryConflict, match="revision conflict"):
        service.create_or_resume(
            document_id=str(document["id"]),
            expected_document_revision=1,
        )
    assert not (tmp_path / "exports" / "document-delivery").exists()


def test_delivery_replays_same_identity_without_rewriting(tmp_path: Path) -> None:
    service, records, _documents, document = _parts(tmp_path)
    first = service.create_or_resume(
        document_id=str(document["id"]), expected_document_revision=1
    )
    stored = records.read("document_deliveries", str(first["delivery_id"]))
    assert stored is not None
    receipt = records.read("document_delivery_receipts", str(first["delivery_id"]))
    assert receipt is not None
    relative = receipt.payload["artifacts"][0]["relative_path"]
    artifact_root = (tmp_path / "exports" / "document-delivery" / relative).parent
    before = {path.name: path.stat().st_mtime_ns for path in artifact_root.iterdir()}

    second = service.create_or_resume(
        document_id=str(document["id"]), expected_document_revision=1
    )

    after = {path.name: path.stat().st_mtime_ns for path in artifact_root.iterdir()}
    assert second["delivery_id"] == first["delivery_id"]
    assert second["replayed"] is True and before == after


def test_uuid_runtime_adopts_pre_uuid_delivery_without_second_authority(
    tmp_path: Path, monkeypatch
) -> None:
    import core.product_core.document_delivery as delivery_module

    service, records, _documents, document = _parts(tmp_path)
    uuid_delivery_id = delivery_module._delivery_id

    def legacy_delivery_id(document_id, revision, formats, style_snapshot):
        value = delivery_module._legacy_delivery_id(
            document_id, revision, formats, style_snapshot
        )
        assert value is not None
        return value

    monkeypatch.setattr(delivery_module, "_delivery_id", legacy_delivery_id)
    legacy = service.create_or_resume(
        document_id=str(document["id"]), expected_document_revision=1
    )
    monkeypatch.setattr(delivery_module, "_delivery_id", uuid_delivery_id)

    replay = service.create_or_resume(
        document_id=str(document["id"]), expected_document_revision=1
    )

    assert replay["delivery_id"] == legacy["delivery_id"]
    assert replay["replayed"] is True
    assert len(records.list("document_deliveries")) == 1
    assert len(records.list("document_delivery_receipts")) == 1
    assert service.artifact(str(legacy["delivery_id"]), "html")[1].startswith(
        b"<!DOCTYPE html>"
    )


def test_delivery_recovers_file_replace_before_terminal_record(tmp_path: Path, monkeypatch) -> None:
    import core.product_core.document_delivery as delivery_module

    service, records, _documents, document = _parts(tmp_path)
    real_publish = delivery_module._publish_exact
    calls = 0

    def crash_after_first_replace(path: Path, data: bytes) -> None:
        nonlocal calls
        real_publish(path, data)
        calls += 1
        if calls == 1:
            raise BaseException("simulated process exit")

    monkeypatch.setattr(delivery_module, "_publish_exact", crash_after_first_replace)
    with pytest.raises(BaseException, match="simulated process exit"):
        service.create_or_resume(
            document_id=str(document["id"]), expected_document_revision=1
        )
    prepared_records = [
        record for record in records.list("document_deliveries")
        if record.payload.get("status") == "prepared"
    ]
    assert len(prepared_records) == 1
    delivery_id = prepared_records[0].object_id
    prepared = records.read("document_deliveries", delivery_id)
    assert prepared is not None and prepared.payload["status"] == "prepared"

    monkeypatch.setattr(delivery_module, "_publish_exact", real_publish)
    _recover_with_core(service)
    recovered = service.projection(delivery_id)
    assert recovered["status"] == "completed"
    assert service.artifact(delivery_id, "markdown")[1].endswith(b"Exact body.")


def test_delivery_reaper_verifies_receipt_saved_before_effect_settle(
    tmp_path: Path, monkeypatch
) -> None:
    service, records, _documents, document = _parts(tmp_path)

    def crash_before_effect_settle(*_args, **_kwargs):
        raise BaseException("simulated process exit before Effect settle")

    monkeypatch.setattr(service._runner, "settle_ok", crash_before_effect_settle)
    with pytest.raises(BaseException, match="before Effect settle"):
        service.create_or_resume(
            document_id=str(document["id"]), expected_document_revision=1
        )

    deliveries = records.list("document_deliveries")
    assert len(deliveries) == 1
    delivery_id = deliveries[0].object_id
    assert deliveries[0].payload["status"] == "prepared"
    assert records.read("document_delivery_receipts", delivery_id) is not None
    assert service._effects.get(delivery_id).state.value == "INFLIGHT"

    outcomes = EffectReaper(service._effects).recover_expired(
        now=int(time.time()) + 301,
        verifiers={
            "document_delivery": lambda effect: service.verify_effect(
                effect.operation_id
            )
        },
    )

    assert [(item.operation_id, item.state.value, item.reason) for item in outcomes] == [
        (delivery_id, "SETTLED_OK", "verifier_resolved")
    ]
    assert service.projection(delivery_id)["status"] == "completed"


def test_delivery_concurrent_first_create_adopts_single_winner(tmp_path: Path) -> None:
    service, records, _documents, document = _parts(tmp_path)

    def create():
        return service.create_or_resume(
            document_id=str(document["id"]), expected_document_revision=1
        )

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = tuple(pool.map(lambda _index: create(), range(2)))

    assert results[0]["delivery_id"] == results[1]["delivery_id"]
    assert sum(result["replayed"] is False for result in results) == 1
    assert len(records.list("document_deliveries")) == 1
    assert len(records.list("document_delivery_receipts")) == 1


def test_new_service_recovers_prepared_with_frozen_historical_renderer(
    tmp_path: Path, monkeypatch
) -> None:
    import core.product_core.document_delivery as delivery_module

    service, records, documents, document = _parts(tmp_path)
    real_publish = delivery_module._publish_exact

    def crash_after_replace(path: Path, data: bytes) -> None:
        real_publish(path, data)
        raise BaseException("simulated process exit")

    monkeypatch.setattr(delivery_module, "_publish_exact", crash_after_replace)
    with pytest.raises(BaseException, match="simulated process exit"):
        service.create_or_resume(
            document_id=str(document["id"]), expected_document_revision=1
        )
    monkeypatch.setattr(delivery_module, "_publish_exact", real_publish)

    style_r1 = {
        "profile_id": "builtin.porcelain-document", "revision": 1,
        "renderer_id": "rebuild.document-html", "renderer_revision": 1,
        "delivery_key": "porcelain-html-v1",
    }
    style_r2 = {
        **style_r1, "revision": 2, "renderer_revision": 2,
        "delivery_key": "porcelain-html-v2",
    }

    class MarkerRenderer:
        def __init__(self, marker: str) -> None:
            self.marker = marker

        def render(self, _document, _markdown):
            return SimpleNamespace(html=self.marker)

    restarted = DocumentDeliveryService(
        SQLiteStructuredRecordStore(records.database_path),
        documents,
        tmp_path,
        style_registry=DocumentDeliveryStyleRegistry(
            current_snapshot=style_r2,
            renderers={
                ("builtin.porcelain-document", 1, "rebuild.document-html", 1): MarkerRenderer("r1"),
                ("builtin.porcelain-document", 2, "rebuild.document-html", 2): MarkerRenderer("r2"),
            },
            delivery_keys={
                ("builtin.porcelain-document", 1, "rebuild.document-html", 1): "porcelain-html-v1",
                ("builtin.porcelain-document", 2, "rebuild.document-html", 2): "porcelain-html-v2",
            },
        ),
    )

    recovery = _recover_with_core(restarted)

    assert recovery["attempted"] == 1 and len(recovery["recovered"]) == 1
    delivery_id = str(recovery["recovered"][0])
    assert restarted.artifact(delivery_id, "html")[1] == b"r1"


def test_distinct_profile_and_renderer_snapshots_use_isolated_artifact_paths(
    tmp_path: Path,
) -> None:
    _service, records, documents, document = _parts(tmp_path)

    class MarkerRenderer:
        def __init__(self, marker: str) -> None:
            self.marker = marker

        def render(self, _document, _markdown):
            return SimpleNamespace(html=self.marker)

    style_a = {
        "profile_id": "profile:alpha", "revision": 1,
        "renderer_id": "renderer:alpha", "renderer_revision": 1,
        "delivery_key": "alpha-v1",
    }
    style_b = {
        "profile_id": "profile-alpha", "revision": 1,
        "renderer_id": "renderer-alpha", "renderer_revision": 1,
        "delivery_key": "beta-v1",
    }

    def service_for(style, marker):
        key = (
            style["profile_id"], style["revision"],
            style["renderer_id"], style["renderer_revision"],
        )
        return DocumentDeliveryService(
            records,
            documents,
            tmp_path,
            style_registry=DocumentDeliveryStyleRegistry(
                current_snapshot=style,
                renderers={key: MarkerRenderer(marker)},
                delivery_keys={key: style["delivery_key"]},
            ),
        )

    first = service_for(style_a, "alpha").create_or_resume(
        document_id=str(document["id"]), expected_document_revision=1, formats=["html"]
    )
    second = service_for(style_b, "beta").create_or_resume(
        document_id=str(document["id"]), expected_document_revision=1, formats=["html"]
    )

    first_record = records.read("document_deliveries", str(first["delivery_id"]))
    second_record = records.read("document_deliveries", str(second["delivery_id"]))
    assert first["delivery_id"] != second["delivery_id"]
    assert first_record is not None and second_record is not None
    first_receipt = records.read("document_delivery_receipts", str(first["delivery_id"]))
    second_receipt = records.read("document_delivery_receipts", str(second["delivery_id"]))
    assert first_receipt is not None and second_receipt is not None
    first_path = first_receipt.payload["artifacts"][0]["relative_path"]
    second_path = second_receipt.payload["artifacts"][0]["relative_path"]
    assert first_path != second_path
    assert service_for(style_a, "alpha").artifact(str(first["delivery_id"]), "html")[1] == b"alpha"
    assert service_for(style_b, "beta").artifact(str(second["delivery_id"]), "html")[1] == b"beta"


def test_long_style_ids_replay_via_short_registered_delivery_key(tmp_path: Path) -> None:
    _service, records, documents, document = _parts(tmp_path)
    profile_id = "p" * 160
    renderer_id = "r" * 160
    style = {
        "profile_id": profile_id, "revision": 1,
        "renderer_id": renderer_id, "renderer_revision": 1,
        "delivery_key": "long-style-v1",
    }
    key = (profile_id, 1, renderer_id, 1)
    service = DocumentDeliveryService(
        records,
        documents,
        tmp_path,
        style_registry=DocumentDeliveryStyleRegistry(
            current_snapshot=style,
            renderers={key: DocumentHtmlRenderer()},
            delivery_keys={key: "long-style-v1"},
        ),
    )

    first = service.create_or_resume(
        document_id=str(document["id"]), expected_document_revision=1
    )
    replay = service.create_or_resume(
        document_id=str(document["id"]), expected_document_revision=1
    )

    assert len(str(first["delivery_id"])) <= 128
    assert replay["delivery_id"] == first["delivery_id"] and replay["replayed"] is True


def test_long_document_id_with_colon_uses_short_sqlite_and_path_identity(tmp_path: Path) -> None:
    _service, records, documents, document = _parts(tmp_path)
    original_id = str(document["id"])
    alias = "document:" + ("long-segment-" * 10)

    class AliasDocuments:
        def read(self, document_id):
            return documents.read(original_id) if document_id == alias else None

        def revision(self, document_id, revision):
            return documents.revision(original_id, revision) if document_id == alias else None

        def markdown(self, document_id, *, revision=None):
            return documents.markdown(original_id, revision=revision) if document_id == alias else None

    service = DocumentDeliveryService(records, AliasDocuments(), tmp_path)
    result = service.create_or_resume(
        document_id=alias, expected_document_revision=1, formats=["docx"]
    )
    stored = records.read("document_deliveries", str(result["delivery_id"]))

    assert stored is not None and len(str(result["delivery_id"])) == 39
    assert stored.payload["document_snapshot"]["document_id"] == alias
    receipt = records.read("document_delivery_receipts", str(result["delivery_id"]))
    assert receipt is not None
    assert ":" not in receipt.payload["artifacts"][0]["relative_path"]
    assert service.artifact(str(result["delivery_id"]), "docx")[0] == "document-r1.docx"


def test_completed_delivery_fails_closed_when_artifact_bytes_drift(tmp_path: Path) -> None:
    service, _records, _documents, document = _parts(tmp_path)
    result = service.create_or_resume(
        document_id=str(document["id"]), expected_document_revision=1
    )
    delivery_id = str(result["delivery_id"])
    path = next(
        (tmp_path / "exports" / "document-delivery").rglob("*.md")
    )
    original = path.read_bytes()
    path.write_bytes(b"x" * len(original))

    with pytest.raises(DocumentDeliveryConflict, match="artifact is unavailable"):
        service.projection(delivery_id)


@pytest.mark.parametrize(
    "formats",
    [[], ["pdf"], ["markdown", "markdown"], "html"],
)
def test_delivery_rejects_unsupported_or_ambiguous_formats(tmp_path: Path, formats) -> None:
    service, _records, _documents, document = _parts(tmp_path)
    with pytest.raises(DocumentDeliveryError, match="formats"):
        service.create_or_resume(
            document_id=str(document["id"]),
            expected_document_revision=1,
            formats=formats,
        )
