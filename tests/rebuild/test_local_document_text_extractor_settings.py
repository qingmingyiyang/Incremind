from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from core.ingestion_core import ObjectStoreSourceRegistrar, SourceSubmission
from core.product_core import (
    AuthorizeLocalDocumentFileForSource,
    GetLocalDocumentTextExtractorSettings,
    LocalDocumentTextExtractorSettingsError,
    RunConfiguredLocalDocumentTextExtractorForSource,
    SaveLocalDocumentTextExtractorSettings,
    ServeLocalDocumentTextExtractorRunEndpoint,
    ServeLocalDocumentTextExtractorSettingsEndpoint,
)
from core.storage_provider import JsonObjectStore


def _store(tmp_path: Path) -> JsonObjectStore:
    return JsonObjectStore(tmp_path / ".rebuild-data", legacy_root=tmp_path / "library")


def _pdf_source(object_store: JsonObjectStore, pdf_path: Path) -> dict[str, object]:
    return ObjectStoreSourceRegistrar(object_store).register(
        SourceSubmission(
            kind="file",
            title="Local document",
            display_name=pdf_path.name,
            media_type="application/pdf",
            size_bytes=pdf_path.stat().st_size,
            file_reference=f"platform-ref-{pdf_path.stem}",
        )
    )


def _docx_source(object_store: JsonObjectStore, docx_path: Path) -> dict[str, object]:
    return ObjectStoreSourceRegistrar(object_store).register(
        SourceSubmission(
            kind="file",
            title="Local DOCX document",
            display_name=docx_path.name,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            size_bytes=docx_path.stat().st_size,
            file_reference=f"platform-ref-{docx_path.stem}",
        )
    )


def test_local_document_text_extractor_settings_are_default_off(tmp_path: Path) -> None:
    object_store = _store(tmp_path)

    settings = GetLocalDocumentTextExtractorSettings(object_store).execute()

    assert settings.enabled is False
    assert settings.status == "disabled"
    assert settings.diagnostic == "disabled_until_explicit_enable"
    assert settings.remote_processing is False
    assert settings.memory_publication == "not_started"


def test_enabling_local_document_text_extractor_requires_confirmation_and_command(
    tmp_path: Path,
) -> None:
    object_store = _store(tmp_path)
    writer = SaveLocalDocumentTextExtractorSettings(object_store)

    try:
        writer.execute(enabled=True, command=(sys.executable,), confirm_enable=False)
    except LocalDocumentTextExtractorSettingsError as error:
        assert str(error) == "enabling local document text extractor requires confirm_enable=true"
    else:
        raise AssertionError("expected explicit enable guard")

    try:
        writer.execute(enabled=True, command=(), confirm_enable=True)
    except LocalDocumentTextExtractorSettingsError as error:
        assert str(error) == "enabled local document text extractor requires command"
    else:
        raise AssertionError("expected command guard")

    settings = writer.execute(
        enabled=True,
        command=(sys.executable, "--version"),
        provider_name="local-python-document-smoke",
        confirm_enable=True,
    )

    assert settings.enabled is True
    assert settings.status == "ready"
    assert settings.provider_name == "local-python-document-smoke"
    assert settings.command == (sys.executable, "--version")


def test_local_document_text_extractor_settings_endpoint_reports_missing_executable(
    tmp_path: Path,
) -> None:
    object_store = _store(tmp_path)
    endpoint = ServeLocalDocumentTextExtractorSettingsEndpoint()
    writer = SaveLocalDocumentTextExtractorSettings(object_store)

    response = endpoint.execute(
        method="PUT",
        path="/api/rebuild/settings/local-document-text-extractor",
        body={
            "enabled": True,
            "command": ["missing-local-document-text-extractor", "{document_path}"],
            "confirm_enable": True,
        },
        get_settings=lambda: GetLocalDocumentTextExtractorSettings(object_store).execute(),
        save_settings=writer.execute,
    )

    assert response.status_code == 200
    assert response.body["status"] == "degraded"
    assert response.body["diagnostic"] == "executable_not_found"
    assert response.body["enabled"] is True


def test_local_document_text_extractor_run_endpoint_rejects_ui_supplied_command(
    tmp_path: Path,
) -> None:
    object_store = _store(tmp_path)
    pdf_path = tmp_path / "document.pdf"
    pdf_path.write_text("PDF 正文", encoding="utf-8")
    source = _pdf_source(object_store, pdf_path)
    endpoint = ServeLocalDocumentTextExtractorRunEndpoint()

    response = endpoint.execute(
        method="POST",
        path=f"/api/rebuild/sources/{source['id']}/document-text",
        body={"command": ["should-not-be-accepted", "{document_path}"]},
        run_extractor=RunConfiguredLocalDocumentTextExtractorForSource(object_store).execute,
    )

    assert response.status_code == 400
    assert response.body["reason"] == (
        "local document text extractor run endpoint does not accept provider command"
    )


def test_configured_local_document_text_extractor_run_uses_saved_settings(
    tmp_path: Path,
) -> None:
    object_store = _store(tmp_path)
    pdf_path = tmp_path / "authorized.pdf"
    pdf_path.write_text("PDF 正文：从设置读取命令。", encoding="utf-8")
    source = _pdf_source(object_store, pdf_path)
    AuthorizeLocalDocumentFileForSource(object_store).execute(
        source_id=str(source["id"]),
        file_path=str(pdf_path),
    )
    SaveLocalDocumentTextExtractorSettings(object_store).execute(
        enabled=True,
        command=(
            sys.executable,
            "-c",
            "from pathlib import Path; import sys; print(Path(sys.argv[1]).read_text(encoding='utf-8'))",
            "{document_path}",
        ),
        provider_name="local-settings-document-text",
        confirm_enable=True,
    )
    endpoint = ServeLocalDocumentTextExtractorRunEndpoint()

    response = endpoint.execute(
        method="POST",
        path=f"/api/rebuild/sources/{source['id']}/document-text",
        body={},
        run_extractor=RunConfiguredLocalDocumentTextExtractorForSource(object_store).execute,
    )
    read_record = object_store.read("source_content_reads", f"content-read-{source['id']}")

    assert response.status_code == 200
    assert response.body["status"] == "completed"
    assert response.body["preview"] == "PDF 正文：从设置读取命令。"
    assert read_record is not None
    assert read_record["text"] == "PDF 正文：从设置读取命令。"


def test_configured_local_document_text_extractor_records_disabled_failure(
    tmp_path: Path,
) -> None:
    object_store = _store(tmp_path)
    pdf_path = tmp_path / "disabled.pdf"
    pdf_path.write_text("disabled body", encoding="utf-8")
    source = _pdf_source(object_store, pdf_path)
    AuthorizeLocalDocumentFileForSource(object_store).execute(
        source_id=str(source["id"]),
        file_path=str(pdf_path),
    )

    result = RunConfiguredLocalDocumentTextExtractorForSource(object_store).execute(
        source_id=str(source["id"])
    )

    assert result.status == "failed"
    assert result.error == "local document text extractor is disabled"


def test_builtin_document_text_extractor_reads_real_docx(tmp_path: Path) -> None:
    object_store = _store(tmp_path)
    document_path = tmp_path / "真实文档.docx"
    document = Document()
    document.add_heading("内建文档提取", level=1)
    document.add_paragraph("DOCX 正文由 bundled python-docx 在本机读取。")
    document.save(str(document_path))
    source = _docx_source(object_store, document_path)
    AuthorizeLocalDocumentFileForSource(object_store).execute(
        source_id=str(source["id"]), file_path=str(document_path)
    )
    SaveLocalDocumentTextExtractorSettings(object_store).execute(
        enabled=True,
        command=("builtin:document-text",),
        provider_name="builtin-document-text",
        confirm_enable=True,
    )

    result = RunConfiguredLocalDocumentTextExtractorForSource(object_store).execute(
        source_id=str(source["id"])
    )

    assert result.status == "completed"
    assert "DOCX 正文由 bundled python-docx 在本机读取" in result.preview


def test_builtin_document_text_extractor_rejects_empty_real_pdf(tmp_path: Path) -> None:
    object_store = _store(tmp_path)
    document_path = tmp_path / "empty.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with document_path.open("wb") as stream:
        writer.write(stream)
    source = _pdf_source(object_store, document_path)
    AuthorizeLocalDocumentFileForSource(object_store).execute(
        source_id=str(source["id"]), file_path=str(document_path)
    )
    SaveLocalDocumentTextExtractorSettings(object_store).execute(
        enabled=True,
        command=("builtin:document-text",),
        provider_name="builtin-document-text",
        confirm_enable=True,
    )

    result = RunConfiguredLocalDocumentTextExtractorForSource(object_store).execute(
        source_id=str(source["id"])
    )

    assert result.status == "failed"
    assert result.error == "local document text extractor returned no text"


def test_builtin_document_text_extractor_reads_real_pdf_text(tmp_path: Path) -> None:
    object_store = _store(tmp_path)
    document_path = tmp_path / "text.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})
    })
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 20 100 Td (Builtin PDF extraction works.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    with document_path.open("wb") as stream:
        writer.write(stream)
    source = _pdf_source(object_store, document_path)
    AuthorizeLocalDocumentFileForSource(object_store).execute(
        source_id=str(source["id"]), file_path=str(document_path)
    )
    SaveLocalDocumentTextExtractorSettings(object_store).execute(
        enabled=True,
        command=("builtin:document-text",),
        provider_name="builtin-document-text",
        confirm_enable=True,
    )

    result = RunConfiguredLocalDocumentTextExtractorForSource(object_store).execute(
        source_id=str(source["id"])
    )

    assert result.status == "completed"
    assert "Builtin PDF extraction works." in result.preview
