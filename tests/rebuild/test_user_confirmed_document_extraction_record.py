from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
SMOKE_SCRIPT = ROOT / "work" / "scripts" / "run_real_document_extraction_smoke.py"
MANUAL_RECORD = ROOT / "docs" / "validation" / "manual-user-confirmed-document-extraction.md"


def test_user_confirmed_docx_smoke_records_confirmation_without_publication(
    tmp_path: Path,
) -> None:
    sample = tmp_path / "user-confirmed-sample.docx"
    document = Document()
    document.add_heading("User confirmed sample", level=1)
    document.add_paragraph("This DOCX stands in for a user-confirmed local document.")
    document.add_paragraph("The extraction must not publish Memory or call external providers.")
    document.save(str(sample))

    completed = subprocess.run(
        [
            sys.executable,
            str(SMOKE_SCRIPT),
            "--output-dir",
            str(tmp_path / "output"),
            "--document-path",
            str(sample),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    output = json.loads(completed.stdout)
    report = json.loads((Path(output["report"])).read_text(encoding="utf-8"))

    assert report["status"] == "passed"
    assert report["mode"] == "user_confirmed_docx"
    assert report["sample"]["user_confirmed"] is True
    assert report["sample"]["path_committed"] is False
    assert report["read_result"]["status"] == "completed"
    assert report["candidate_status"] == "pending_review"
    assert report["candidate_auto_promote_allowed"] is False
    assert report["memory_atoms_written"] == 0
    assert report["staging_atoms_written"] == 0
    assert report["memory_publications_written"] == 0
    assert report["privacy_boundary"]["user_private_file_read"] is True
    assert report["privacy_boundary"]["full_document_text_committed"] is False
    assert report["privacy_boundary"]["external_model_provider_called"] is False
    assert report["privacy_boundary"]["api_key_committed"] is False
    assert report["privacy_boundary"]["os_path_in_source_contract"] is False


def test_user_confirmed_document_manual_record_keeps_sensitive_boundaries() -> None:
    record = MANUAL_RECORD.read_text(encoding="utf-8")

    assert "user_confirmed_docx" in record
    assert "产品设计信息提取文档_个人AI记忆工作台.docx" in record
    assert "path_committed: false" in record
    assert "full_document_text_committed: false" in record
    assert "external_model_provider_called: false" in record
    assert "api_key_committed: false" in record
    assert "memory_atoms_written: 0" in record
    assert "staging_atoms_written: 0" in record
    assert "memory_publications_written: 0" in record
    assert "pending_review" in record
    assert "D:\\Users\\Chrip\\Downloads" not in record
    assert "sk-" not in record
